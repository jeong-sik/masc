"""Exercise the embedded parser against real, locally generated PPTX packages.

Run with the managed presentation interpreter; no network or renderer required.
"""

import ast
import io
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

SOURCE = (
    Path(__file__).resolve().parents[2]
    / "lib/verification_presentation_inspection_parser.ml"
).read_text()
PYTHON = SOURCE.split("{python|", 1)[1].rsplit("|python}", 1)[0]
NAMESPACE = {}
exec(compile(PYTHON.rsplit("\nmain()", 1)[0], "presentation-parser", "exec"), NAMESPACE)


class PresentationParserTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / "source.pptx"
        deck = Presentation()
        for index in range(2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text = "Slide " + str(index + 1)
            if index == 0:
                box.text_frame.paragraphs[0].runs[
                    0
                ].hyperlink.address = "https://example.invalid/reference"
            else:
                slide._element.set("show", "0")
        deck.save(str(self.source))

    def inspect(self):
        return NAMESPACE["inspect"](self.source, sys.prefix)

    def test_visibility_and_hyperlink_ownership(self):
        slides = self.inspect()["slides"]
        self.assertEqual([True, False], [slide["visible"] for slide in slides])
        self.assertEqual(["https://example.invalid/reference"], slides[0]["hyperlinks"])
        self.assertEqual([], slides[1]["hyperlinks"])

    def test_python38_syntax_and_containment(self):
        ast.parse(PYTHON, feature_version=(3, 8))
        # Python 3.8 does not provide is_relative_to. Exercise the fallback itself.
        with patch.object(
            Path,
            "is_relative_to",
            side_effect=AssertionError("Python 3.9 API"),
            create=True,
        ):
            self.assertTrue(self.inspect()["ok"])

    def test_zip_bomb_rejected_before_crc_decompression(self):
        with zipfile.ZipFile(
            self.source, "w", compression=zipfile.ZIP_DEFLATED
        ) as archive:
            archive.writestr("bomb.xml", b"0" * (2 * 1024 * 1024))
        with patch.object(
            zipfile.ZipFile,
            "testzip",
            side_effect=AssertionError("decompressed before preflight"),
        ):
            with self.assertRaises(NAMESPACE["PolicyError"]):
                self.inspect()

    def test_entry_count_rejected_before_crc_decompression(self):
        with zipfile.ZipFile(self.source, "w") as archive:
            for index in range(4097):
                archive.writestr(str(index), b"")
        with patch.object(
            zipfile.ZipFile,
            "testzip",
            side_effect=AssertionError("decompressed before preflight"),
        ):
            with self.assertRaises(NAMESPACE["PolicyError"]):
                self.inspect()

    def test_size_preflight_rejects_member_and_total(self):
        for sizes in ([17 * 1024 * 1024], [16 * 1024 * 1024] * 9):
            entries = []
            for index, size in enumerate(sizes):
                entry = zipfile.ZipInfo(str(index))
                entry.file_size = size
                entry.compress_size = size
                entries.append(entry)
            with (
                self.subTest(sizes=sizes),
                patch.object(zipfile.ZipFile, "infolist", return_value=entries),
                patch.object(
                    zipfile.ZipFile,
                    "testzip",
                    side_effect=AssertionError("decompressed before preflight"),
                ),
            ):
                with self.assertRaises(NAMESPACE["PolicyError"]):
                    self.inspect()

    def test_external_loading_is_rejected(self):
        original = self.source.read_bytes()
        from xml.etree import ElementTree as etree

        with (
            zipfile.ZipFile(io.BytesIO(original)) as source,
            zipfile.ZipFile(self.source, "w") as target,
        ):
            for entry in source.infolist():
                data = source.read(entry)
                if entry.filename == "ppt/slides/_rels/slide1.xml.rels":
                    document = etree.fromstring(data)
                    for relation in document:
                        if relation.get("TargetMode") == "External":
                            relation.set(
                                "Type",
                                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                            )
                    data = etree.tostring(document)
                target.writestr(entry, data)
        with self.assertRaises(NAMESPACE["PolicyError"]):
            self.inspect()


if __name__ == "__main__":
    unittest.main()
