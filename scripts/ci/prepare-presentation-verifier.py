"""Create a managed parser and actual PPTX inputs for the native CI feature test.

Run explicitly by CI; product Read never installs software. All files belong to
the runner's temporary directory, including the python-pptx installation.
"""
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile


def main():
    contract_spec = importlib.util.spec_from_file_location(
        "presentation_fixture", Path(__file__).with_name("presentation_fixture.py"))
    contract = importlib.util.module_from_spec(contract_spec)
    contract_spec.loader.exec_module(contract)
    contract.descriptor_path()  # Fail clearly outside CI before installing anything.
    if len(sys.argv) == 1:
        # RUNNER_TEMP can be inside HOME; only the descriptor belongs there.
        base = (Path(tempfile.gettempdir()) / "masc-presentation-verifier").resolve()
    elif len(sys.argv) == 3 and sys.argv[1] == "generate":
        base = Path(sys.argv[2])
    else:
        raise SystemExit("expected no arguments, or generate BASE under the managed interpreter")
    if not base.is_absolute() or base.resolve().is_relative_to(Path.home().resolve()):
        raise SystemExit("presentation fixture workspace must be absolute and outside HOME")
    environment = base / ".masc/runtime-tools/presentation"
    python = environment / "bin/python3"
    if len(sys.argv) == 1:
        subprocess.run([sys.executable, "-I", "-m", "venv", "--copies", str(environment)], check=True)
        subprocess.run([str(python), "-I", "-m", "pip", "--isolated", "install",
                        "--require-virtualenv", "--only-binary=:all:", "python-pptx"], check=True)
        subprocess.run([str(python), "-I", "-B", str(Path(__file__).resolve()), "generate", str(base)], check=True)
        contract.publish(base)
        return
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    for title, note in [("First: leave a memory", "Speaker note one"),
                        ("Second: keep your choice", "Speaker note two")]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = title
        slide.notes_slide.notes_text_frame.text = note
    # Hidden slides must still be represented in the independent inspection.
    deck.slides[1]._element.set("show", "0")
    table = deck.slides[1].shapes.add_table(1, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Leave"
    table.cell(0, 1).text = "Take"
    # Normal clickable references should not require the renderer to fetch a URL.
    paragraph = deck.slides[0].shapes[0].text_frame.paragraphs[0]
    paragraph.runs[0].hyperlink.address = "https://example.invalid/reference"
    fixture = base / "inputs"
    fixture.mkdir(parents=True, exist_ok=True)
    source = fixture / "presentation.pptx"
    deck.save(source)
    original = source.read_bytes()
    (fixture / "expected.json").write_text(json.dumps({
        "bytes": len(original), "sha256": hashlib.sha256(original).hexdigest(),
        "titles": ["First: leave a memory", "Second: keep your choice"],
        "notes": ["Speaker note one", "Speaker note two"],
    }) + "\n")
    (fixture / "broken.pptx").write_bytes(original[:len(original) // 2])
    # An auto-loaded external image is a different relationship from a hyperlink.
    from lxml import etree
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(fixture / "external.pptx", "w") as dst:
        for entry in src.infolist():
            data = src.read(entry)
            if entry.filename == "ppt/slides/_rels/slide1.xml.rels":
                root = etree.fromstring(data)
                etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="external-image", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                    Target="https://example.invalid/image.png", TargetMode="External")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
            dst.writestr(entry, data)
    print(json.dumps({"base_path": str(base), "fixture_sha256": hashlib.sha256(original).hexdigest()}))


if __name__ == "__main__":
    main()
